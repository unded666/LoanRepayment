from fastapi import FastAPI, Request, Form, UploadFile, File
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from finance import (
    AmortizationComparisonDifference,
    AmortizationComparisonResult,
    calculate_amortization
)
from datetime import date
from pydantic import BaseModel
import requests
import io
import base64
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from PIL import Image as PILImage

app = FastAPI()
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

class CalculationResponseWithCurrency(BaseModel):
    result: AmortizationComparisonResult
    currency_symbol: str

COUNTRY_CURRENCY = {
    'US': '$', 'CA': '$', 'GB': '£', 'DE': '€', 'FR': '€', 'ES': '€', 'IT': '€', 'IE': '€',
    'JP': '¥', 'CN': '¥', 'IN': '₹', 'AU': '$', 'NZ': '$', 'SG': '$', 'ZA': 'R', 'CH': 'CHF',
    'SE': 'kr', 'NO': 'kr', 'DK': 'kr', 'PL': 'zł', 'CZ': 'Kč', 'RU': '₽', 'BR': 'R$', 'MX': '$',
    'KR': '₩', 'TR': '₺', 'IL': '₪', 'SA': '﷼', 'AE': 'د.إ', 'HK': '$', 'MY': 'RM', 'TH': '฿',
    'ID': 'Rp', 'PH': '₱', 'NG': '₦', 'EG': '£', 'PK': '₨', 'BD': '৳', 'UA': '₴', 'AR': '$',
    'CL': '$', 'CO': '$', 'PE': 'S/', 'VE': 'Bs', 'VN': '₫', 'TW': 'NT$', 'HU': 'Ft',
}

def get_currency_symbol_from_ip(ip: str) -> str:
    try:
        resp = requests.get(f"https://ipapi.co/{ip}/json/", timeout=2)
        if resp.status_code == 200:
            data = resp.json()
            country = data.get('country')
            if country and country in COUNTRY_CURRENCY:
                return COUNTRY_CURRENCY[country]
    except Exception:
        pass
    return '$'  # Default to USD

@app.get("/", response_class=HTMLResponse)
async def read_root(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})

@app.post("/calculate", response_model=CalculationResponseWithCurrency)
async def calculate(
    request: Request,
    purchase_price: float = Form(...),
    interest_rate: float = Form(...),
    down_payment: float = Form(...),
    loan_term: int = Form(...),
    start_date: str = Form(...),
    custom_repayment: float = Form(None)
):
    client_ip = request.client.host if request.client else None
    if client_ip in ("127.0.0.1", "::1", None):
        xff = request.headers.get("x-forwarded-for")
        if xff:
            client_ip = xff.split(",")[0].strip()
    currency_symbol = get_currency_symbol_from_ip(client_ip or "")
    orig_result = calculate_amortization(
        purchase_price=purchase_price,
        interest_rate=interest_rate,
        down_payment=down_payment,
        loan_term_years=loan_term,
        start_date=date.fromisoformat(start_date)
    )
    custom_result = None
    diff = None
    if custom_repayment and custom_repayment > orig_result.summary.monthly_payment:
        custom_result = calculate_amortization(
            purchase_price=purchase_price,
            interest_rate=interest_rate,
            down_payment=down_payment,
            loan_term_years=loan_term,
            start_date=date.fromisoformat(start_date),
            custom_repayment=custom_repayment
        )
        diff = AmortizationComparisonDifference(
            total_payments_diff=round(orig_result.summary.total_payments - custom_result.summary.total_payments, 2),
            months_diff=len(orig_result.schedule) - len(custom_result.schedule)
        )
    result = AmortizationComparisonResult(
        original=orig_result,
        custom=custom_result,
        difference=diff
    )
    return CalculationResponseWithCurrency(result=result, currency_symbol=currency_symbol)

@app.post("/download_excel")
async def download_excel(
    request: Request,
    purchase_price: float = Form(...),
    interest_rate: float = Form(...),
    down_payment: float = Form(...),
    loan_term: int = Form(...),
    start_date: str = Form(...),
    custom_repayment: float = Form(None),
    chart1_base64: str = Form(...),
    chart2_base64: str = Form(...)
):
    # Calculate results
    orig_result = calculate_amortization(
        purchase_price=purchase_price,
        interest_rate=interest_rate,
        down_payment=down_payment,
        loan_term_years=loan_term,
        start_date=date.fromisoformat(start_date)
    )
    custom_result = None
    diff = None
    if custom_repayment and custom_repayment > orig_result.summary.monthly_payment:
        custom_result = calculate_amortization(
            purchase_price=purchase_price,
            interest_rate=interest_rate,
            down_payment=down_payment,
            loan_term_years=loan_term,
            start_date=date.fromisoformat(start_date),
            custom_repayment=custom_repayment
        )
        diff = AmortizationComparisonDifference(
            total_payments_diff=round(orig_result.summary.total_payments - custom_result.summary.total_payments, 2),
            months_diff=len(orig_result.schedule) - len(custom_result.schedule)
        )
    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_schedule = wb.create_sheet("Schedule")
    # Write summary
    ws_summary.append(["Monthly Repayment", orig_result.summary.monthly_payment])
    ws_summary.append(["Total Payments", orig_result.summary.total_payments])
    ws_summary.append(["Total Interest", orig_result.summary.total_interest])
    ws_summary.append(["Payoff Date", str(orig_result.summary.payoff_date)])
    if diff:
        ws_summary.append(["--- With Custom Repayment ---"])
        ws_summary.append(["Monthly Repayment", custom_result.summary.monthly_payment])
        ws_summary.append(["Total Payments", custom_result.summary.total_payments])
        ws_summary.append(["Total Interest", custom_result.summary.total_interest])
        ws_summary.append(["Payoff Date", str(custom_result.summary.payoff_date)])
        ws_summary.append(["Months Saved", diff.months_diff])
        ws_summary.append(["Total Payments Saved", diff.total_payments_diff])
    # Add chart images
    def add_chart(ws, chart_base64, cell):
        if chart_base64:
            img_bytes = base64.b64decode(chart_base64.split(",")[-1])
            img = PILImage.open(io.BytesIO(img_bytes))
            img_path = f"/tmp/chart_{cell}.png"
            img.save(img_path)
            ws.add_image(OpenpyxlImage(img_path), cell)
    add_chart(ws_summary, chart1_base64, "A10")
    add_chart(ws_summary, chart2_base64, "A30")
    # Write schedule
    ws_schedule.append(["#", "Date", "Payment", "Principal", "Interest", "Balance"])
    for row in orig_result.schedule:
        ws_schedule.append([
            row.payment_number, str(row.date), row.payment, row.principal, row.interest, row.balance
        ])
    # Save to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=amortization.xlsx"}
    )

def main():
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)

if __name__ == "__main__":
    main()
